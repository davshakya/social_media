"""Editable vertical PowerPoint decks generated from validated storyboards."""
from pathlib import Path


def create_presentation(story, destination):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError as exc:
        raise RuntimeError('Install PowerPoint support: .\\.venv\\Scripts\\python.exe -m pip install -e ".[ppt]"') from exc

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(7.5), Inches(13.333)
    blank = deck.slide_layouts[6]
    colors = [(21, 33, 68), (23, 78, 110), (75, 44, 128), (178, 66, 92), (13, 128, 127)]

    def textbox(slide, text, left, top, width, height, size, color=(255, 255, 255), bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text, paragraph.alignment = text, align
        run = paragraph.runs[0]
        run.font.name, run.font.size, run.font.bold = "Aptos Display", Pt(size), bold
        run.font.color.rgb = RGBColor(*color)
        return box

    for index, scene in enumerate(story.scenes):
        slide = deck.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(*colors[index % len(colors)])
        accent = (71, 235, 214) if index % 2 == 0 else (255, 205, 104)
        accent_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.42), Inches(.5), Inches(6.66), Inches(.38))
        accent_shape.fill.solid()
        accent_shape.fill.fore_color.rgb = RGBColor(*accent)
        accent_shape.line.fill.background()
        textbox(slide, "TECHGYAAN  •  QUICK EXPLAINER", .5, .55, 6.5, .35, 13, accent, True, PP_ALIGN.CENTER)
        textbox(slide, story.topic, .55, 1.35, 6.4, 1.25, 29, bold=True, align=PP_ALIGN.CENTER)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.55), Inches(3.0), Inches(6.4), Inches(5.6))
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(10, 19, 42)
        card.line.color.rgb = RGBColor(*accent)
        textbox(slide, f"{index + 1:02d}  {scene.caption.upper()}", .85, 3.45, 5.8, .6, 18, accent, True, PP_ALIGN.CENTER)
        textbox(slide, scene.narration, .92, 4.4, 5.65, 2.7, 22, align=PP_ALIGN.CENTER)
        textbox(slide, scene.action.replace("_", " ").upper(), .85, 7.55, 5.8, .45, 13, (190, 205, 230), True, PP_ALIGN.CENTER)
        textbox(slide, story.cta, .6, 10.3, 6.3, .65, 18, (240, 246, 255), True, PP_ALIGN.CENTER)
        textbox(slide, f"SLIDE {index + 1} OF {len(story.scenes)}", .6, 11.8, 6.3, .35, 12, (190, 205, 230), align=PP_ALIGN.CENTER)
    deck.save(destination)
    return destination
